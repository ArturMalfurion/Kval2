# -*- coding: utf-8 -*-
"""Синхронизация KB (Bitrix24) с динамикой продаж 1С (Din_RRS / Din_TP)
и квалификация сделок. Результат: Output.xlsx (листы Результат и Сводка).
Спецификация: AGENTS.md"""
import os
import re
import math
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

FOLDER = os.path.dirname(os.path.abspath(__file__))
INPUT_DIR = os.path.join(FOLDER, "input")
OUTPUT_DIR = os.path.join(FOLDER, "output")
KB_FILE = os.path.join(INPUT_DIR, "KB.xlsx")
RRS_FILE = os.path.join(INPUT_DIR, "Din_RRS.xlsx")
TP_FILE = os.path.join(INPUT_DIR, "Din_TP.xlsx")
OUT_FILE = os.path.join(OUTPUT_DIR, "Output.xlsx")

LAT2CYR = str.maketrans("COAEPHKTMXBY", "СУАЕРНКТМХБУ")
ORG_FORMS = r"\b(ООО|ИП|АО|ПАО|ОАО|МУП|ГУП|УП|НПО|СЗ|ССК|ЧТПУП|ЕП|ЗАО)\b"
CITY_STOP = r"\b(САМАРА|МОСКВА|ЖИГУЛЕВСК|КАЗАНЬ|КАЗАНЬЮ|КАЗАНИ|ЗЕЛЕНОДОЛЬСК|ЗЕЛЕНОДОЛЬСКУ|АЗИНО|АЗИНИ|КВАРТАЛ|КВАРТАЛА|КВАРТАЛУ|ВОЛЖСК|ВОЛЖСКА|АЛЕКСЕЕВСКОЕ|АЛЕКСЕЕВСКОГО|НОВОНИКОЛАЕВСКИЙ|НОВОНИКОЛАЕВСКИМ)\b"
GEN_STOP = {"ТРАНС", "СЕРВИС", "СТРОЙ", "САМАРА", "АВТО", "КАЗАНЬ", "ЗЕЛЕНОДОЛЬСК",
            "АЗИНО", "КВАРТАЛ", "КВАРТАЛА", "ВОЛЖСК", "АЛЕКСЕЕВСКОЕ"}
TECH = [r"\b\d+\s*ДН\b", r"\bДН\b", r"ЭДО", r"Э\s*ДО\b", r"СКАН", r"\bQR\b",
        r"БЕЗ\s+ДОГОВОРА", r"Б\s*/\s*Д", r"РАСПРЕДЕЛЯТЬ\s+ОПЛАТЫ", r"РАСПРЕДЕЛЯТЬ",
        r"РАСПРЕДелять", r"\+\+\++", r"САЖАТЬ\s+НА\s+АК\.?\d*", r"НЕ\s+ПЛАТИМ",
        r"НОВЫЙ\s+ДОГОВОР", r"ПРЕДОПЛАТА", r"Б\s*/\s*Д"]

def normalize(text):
    if text is None or (isinstance(text, float) and math.isnan(text)):
        return ""
    s = str(text).replace("\xa0", " ")
    s = s.upper().translate(LAT2CYR)
    s = re.sub(r"\([^)]*\)", " ", s)
    s = re.sub(r"\[[^\]]*\]", " ", s)
    s = re.sub(r"[«»\"'\u2018\u2019]", " ", s)
    s = re.sub(r"ТОП\s*-", " ", s)
    s = re.sub(r"\+?7[\s(\-]*\d{3}[\s)\-]*\d{3}[\s\-]*\d{2}[\s\-]*\d{2}", " ", s)
    s = re.sub(r"\b8?9\d{9}\b", " ", s)
    s = re.sub(r"\b\d{7,}\b", " ", s)
    s = re.sub(r"\b(УЛ|УЛИЦА|ПРОЕЗД|ШОССЕ|ЛИТЕР|КОРПУС|ОФИС|ПОС|СНТ|КМ|ТП|ЖД|СТ)\b\.?", " ", s)
    s = re.sub(r"\bГ\.\s*[А-ЯЁ][А-ЯЁ\-]+", " ", s)
    s = re.sub(r"\bД\.?\s*\d+[А-ЯЁa-z]?", " ", s)
    s = re.sub(r"\bС\.\s*[А-ЯЁ][А-ЯЁ\-]+", " ", s)
    s = re.sub(r"\d+\s*[()]\s*\d+\s*ДН", " ", s)
    for pat in TECH:
        s = re.sub(pat, " ", s)
    s = re.sub(ORG_FORMS, " ", s)
    s = re.sub(CITY_STOP, " ", s)
    s = re.sub(r"[^А-ЯЁA-Z0-9 ]", " ", s)
    s = re.sub(r"\b\d+\b", " ", s)
    s = re.sub(r"\s+", " ", s).strip()
    return s

def parse_potential(text):
    if text is None or (isinstance(text, float) and math.isnan(text)):
        return 0.0
    s = str(text).replace("\xa0", " ").replace(" ", " ")
    m = re.search(r"(\d[\d ]*(?:[.,]\d+)?)", s)
    if not m:
        return 0.0
    num = m.group(1).replace(" ", "").replace(",", ".")
    try:
        return float(num)
    except ValueError:
        return 0.0

def norm_to_num(val):
    if val is None or (isinstance(val, float) and math.isnan(val)):
        return 0.0
    if isinstance(val, (int, float)):
        return float(val)
    s = str(val).replace("\xa0", " ").strip()
    if not s or s in ("-", "—"):
        return 0.0
    m = re.search(r"-?\d[\d ]*(?:[.,]\d+)?", s)
    if not m:
        return 0.0
    try:
        return float(m.group(0).replace(" ", "").replace(",", "."))
    except ValueError:
        return 0.0
PATR = re.compile(r"(ИЧ|ОВНА|ЕВНА|ИЧНА|ЫНЫ|КЗЫ|ОГЛЫ)$")
SURNAME = re.compile(r"(ОВ|ЕВ|ИВ|ЫХ|ИХ|АН|ЯН|ЕНКО|СКИ|ЦКИЙ|ШВИЛИ|УЛЫ|ОГЛЫ|ДЗЕ)$")

def canon_lead(v):
    v = str(v).strip().replace("\xa0", " ").upper()
    if v.startswith("ТАТПАРТ"):
        return "ТатПартс"
    if v.startswith("РАРОС"):
        return "Рарос"
    return str(v).strip()

def is_person(text):
    s = str(text)
    if re.search(r"\bИП\b", s.upper().translate(LAT2CYR)):
        return True
    toks = normalize(text).split()
    if not 2 <= len(toks) <= 4 or not toks[0]:
        return False
    rest = toks[1:]
    if any(PATR.search(w) for w in rest):
        return True
    if SURNAME.search(toks[0]) and all(len(w) <= 2 for w in rest):
        return True
    return False

def gen_tokens(norm_str):
    toks = set(norm_str.split())
    return {t for t in toks if len(t) > 3 and t not in GEN_STOP}

ABBR = re.compile(r"\b[А-ЯЁ]{2,5}\b")
ABBR_EXCL = {"ООО", "ИП", "АО", "ПАО", "ОАО", "МУП", "ГУП", "УП", "НПО", "СЗ", "ССК",
             "ТОП", "РФ", "ТД", "ПК", "СК", "ТП", "ПО", "ДА", "ГК", "МТК", "КПД", "СНТ", "ТК"}
def get_abbrs(raw):
    """Аббревиатуры: только реально ЗАГЛАВНЫЕ токены исходной строки KB."""
    s = str(raw).replace("\xa0", " ")
    res = set()
    for w in ABBR.findall(s):
        if w in ABBR_EXCL or w in GEN_STOP:
            continue
        res.add(w)
    return res

def has_abbr(ab, raw_up):
    return re.search(r"(?<![А-ЯЁA-Z])" + ab + r"(?![А-ЯЁA-Z])", raw_up) is not None

STRUCT_WORDS = {"КЛИЕНТ", "НАПРАВЛЕНИЕ", "САМАРА", "ТОЛЬЯТТИ", "СЫЗРАНЬ", "ВОСТОК", "ЗАПАД"}

def is_structural(name):
    toks = set(normalize(name).split())
    return bool(toks) and toks <= STRUCT_WORDS

def load_dynamics(path):
    df = pd.read_excel(path, sheet_name=0)
    df = df.rename(columns={df.columns[0]: "Клиент"})
    month_cols = [c for c in df.columns[1:4]]
    df = df.dropna(subset=["Клиент"]).copy()
    df["Клиент"] = df["Клиент"].astype(str).str.strip()
    df = df[df["Клиент"] != ""]
    df = df[~df["Клиент"].map(is_structural)].reset_index(drop=True)
    for c in month_cols:
        df[c] = df[c].map(norm_to_num)
    df["avg"] = df[month_cols].sum(axis=1) / 3.0
    df["norm"] = df["Клиент"].map(normalize)
    df["raw_up"] = df["Клиент"].str.upper().str.replace("\xa0", " ", regex=False)
    return df, month_cols

def match_one(kb_name, df, used):
    n = normalize(kb_name)
    if not n:
        return None
    cand = df.index[~df.index.isin(used)].tolist()
    sub = df.loc[cand]
    # 1. Exact (множество токенов)
    hits = sub[sub["norm"] != ""]
    hits = hits[hits.apply(lambda r: set(r["norm"].split()) == set(n.split()), axis=1)]
    if len(hits):
        return int(hits.index[0])
    # 2. Core
    if not is_person(kb_name):
        kb_core = gen_tokens(n)
        if kb_core:
            best, best_len = None, 0
            for idx, r in sub.iterrows():
                if not r["norm"]:
                    continue
                cl_core = gen_tokens(r["norm"])
                if not cl_core:
                    continue
                inter = kb_core & cl_core
                if inter and (kb_core <= cl_core or cl_core <= kb_core):
                    if len(inter) > best_len:
                        best, best_len = idx, len(inter)
            if best is not None:
                return int(best)
    # 3. Abbr
    abbrs = get_abbrs(kb_name)
    if abbrs:
        best, best_key = None, None
        for idx, r in sub.iterrows():
            matched = [a for a in abbrs if has_abbr(a, r["raw_up"])]
            if matched:
                key = (len(matched), max(map(len, matched)))
                if best_key is None or key > best_key:
                    best, best_key = idx, key
        if best is not None:
            return int(best)
    # 4. IP / ФИО
    toks = n.split()
    if is_person(kb_name) and len(toks) >= 2:
        surname, ini = toks[0], [w[0] for w in toks[1:] if w]
        best = None
        for idx, r in sub.iterrows():
            rt = r["norm"].split()
            if len(rt) < 1 or rt[0] != surname:
                continue
            rini = [w[0] for w in rt[1:] if w]
            k = min(len(ini), len(rini))
            if k >= 1 and ini[:k] == rini[:k]:
                best = idx
                break
        if best is not None:
            return int(best)
    return None

def is_region_col(header):
    return re.sub(r"[^а-яёa-z]", "", str(header).lower()) == "бизнесрегион"

HDR_MAP = {
    "названиесделки": "Название сделки",
    "типлида": "Тип лида",
    "стадиясделки": "Стадия сделки",
    "потенциалзакупокпорароситатпартс": "Потенциал закупок по Рарос и ТатПартс",
}

def main():
    kb_raw = pd.read_excel(KB_FILE, sheet_name=0)
    canon = {}
    for c in kb_raw.columns:
        key = re.sub(r"[^а-яёa-z]", "", str(c).lower())
        if key in HDR_MAP:
            canon[HDR_MAP[key]] = c
    region_col = next((c for c in kb_raw.columns if is_region_col(c)), None)
    missing = [v for v in ("Название сделки", "Тип лида", "Стадия сделки",
                           "Потенциал закупок по Рарос и ТатПартс") if v not in canon]
    if missing:
        raise RuntimeError(f"В KB.xlsx не найдены столбцы: {missing}")
    kb = kb_raw[[canon["Название сделки"], canon["Тип лида"], canon["Стадия сделки"],
                 canon["Потенциал закупок по Рарос и ТатПартс"]]].copy()
    kb.columns = ["Название сделки", "Тип лида", "Стадия сделки",
                  "Потенциал закупок по Рарос и ТатПартс"]
    kb["Бизнес-регион"] = (kb_raw[region_col].fillna("").astype(str).str.strip()
                           if region_col else "")
    kb = kb[["Название сделки", "Тип лида", "Стадия сделки",
             "Потенциал закупок по Рарос и ТатПартс", "Бизнес-регион"]]
    kb = kb.dropna(subset=["Название сделки"]).reset_index(drop=True)
    kb["Тип лида"] = kb["Тип лида"].fillna("").astype(str).map(canon_lead)
    kb["Потенциал"] = kb["Потенциал закупок по Рарос и ТатПартс"].map(parse_potential)

    rrs, m_rrs = load_dynamics(RRS_FILE)
    tp, m_tp = load_dynamics(TP_FILE)
    print(f"KB строк: {len(kb)}; Din_RRS клиентов: {len(rrs)} {m_rrs}; Din_TP клиентов: {len(tp)} {m_tp}")

    pairs = kb.groupby(kb["Название сделки"].astype(str).str.strip())["Тип лида"] \
              .apply(lambda s: {"Рарос", "ТатПартс"} <= set(s))

    used_rrs, used_tp = set(), set()
    match_cache = {}
    results = []
    for i, row in kb.iterrows():
        name = str(row["Название сделки"]).strip()
        lead = row["Тип лида"]
        pot = row["Потенциал"]
        rec_stage = row["Стадия сделки"]
        if lead == "Рарос":
            df, used = rrs, used_rrs
        elif lead == "ТатПартс":
            df, used = tp, used_tp
        else:
            df, used = None, None
        client, avg = "Не найдено", 0.0
        if df is not None:
            key = (lead, name)
            if key in match_cache:
                idx = match_cache[key]
            else:
                idx = match_one(name, df, used)
                if idx is not None:
                    used.add(idx)
                match_cache[key] = idx
            if idx is not None:
                client = df.at[idx, "Клиент"]
                avg = float(df.at[idx, "avg"])
        found = df is not None and client != "Не найдено"
        proj = pot / 2 if pairs.get(name, False) else pot
        if lead == "":
            rec = "Ошибка типа лида"
        elif not found:
            rec = ("Проверьте первый заказ и заполните потенциал" if pot < 1000
                   else "Проверьте первый заказ")
        else:
            if pot < 1000:
                rec = "Ошибка, заполните потенциал"
            elif avg == 0:
                rec = "Отвалился"
            elif avg >= proj:
                rec = "Есть выборка потенциала"
            else:
                rec = "Нет выборки потенциала"
        results.append({
            "Название сделки": name,
            "Тип лида": lead,
            "Клиент в динамике": client,
            "Стадия сделки": rec_stage,
            "Потенциал закупок по Рарос и ТатПартс": pot,
            "Средняя выручка за квартал": avg,
            "Рекомендуемая стадия": rec,
            "Бизнес-регион": row["Бизнес-регион"],
        })
    out = pd.DataFrame(results)
    order = {"Рарос": 0, "ТатПартс": 1, "": 2}
    out["_o"] = out["Тип лида"].map(order)
    out = out.sort_values(["_o", "Название сделки"]).drop(columns="_o").reset_index(drop=True)

    # ---- Output.xlsx ----
    wb = Workbook()
    ws = wb.active
    ws.title = "Результат"
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="305496")
    fill_by_stage = {
        "Ошибка типа лида": "FFC7CE",
        "Отвалился": "FFC7CE",
        "Ошибка, заполните потенциал": "FFEB9C",
        "Проверьте первый заказ и заполните потенциал": "FFEB9C",
        "Нет выборки потенциала": "DDEBF7",
        "Есть выборка потенциала": "C6EFCE",
        "Проверьте первый заказ": "DDEBF7",
    }
    for c, col in enumerate(out.columns, 1):
        cell = ws.cell(1, c, col)
        cell.font, cell.fill = head_font, head_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for r, (_, rw) in enumerate(out.iterrows(), 2):
        ws.cell(r, 1, rw["Название сделки"])
        ws.cell(r, 2, rw["Тип лида"])
        ws.cell(r, 3, rw["Клиент в динамике"])
        ws.cell(r, 4, rw["Стадия сделки"])
        ws.cell(r, 5, rw["Потенциал закупок по Рарос и ТатПартс"]).number_format = "# ##0,00"
        ws.cell(r, 6, rw["Средняя выручка за квартал"]).number_format = "# ##0,00"
        st = ws.cell(r, 7, rw["Рекомендуемая стадия"])
        if rw["Рекомендуемая стадия"] in fill_by_stage:
            st.fill = PatternFill("solid", fgColor=fill_by_stage[rw["Рекомендуемая стадия"]])
        ws.cell(r, 8, rw["Бизнес-регион"])
    for c, w in zip(range(1, 9), (38, 12, 38, 22, 20, 20, 24, 18)):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = "A2"

    ws2 = wb.create_sheet("Сводка")
    total = len(out)
    ok = int((out["Клиент в динамике"] != "Не найдено").sum())
    nf = total - ok
    ws2.cell(1, 1, "Показатель").font = Font(bold=True)
    ws2.cell(1, 2, "Значение").font = Font(bold=True)
    rows = [("Всего строк", total),
            ("Синхронизировано успешно", f"{ok} ({ok/total:.1%})"),
            ("Не найдено", f"{nf} ({nf/total:.1%})")]
    for stg, cnt in out["Рекомендуемая стадия"].value_counts().items():
        rows.append((stg, int(cnt)))
    for r, (k, v) in enumerate(rows, 2):
        ws2.cell(r, 1, k)
        ws2.cell(r, 2, v)
    ws2.column_dimensions["A"].width = 45
    ws2.column_dimensions["B"].width = 16
    os.makedirs(OUTPUT_DIR, exist_ok=True)
    wb.save(OUT_FILE)

    pd.set_option("display.max_rows", None)
    pd.set_option("display.max_columns", None)
    pd.set_option("display.width", 250)
    pd.set_option("display.max_colwidth", 60)
    print(out.to_string())
    print("\n===== СВОДКА =====")
    for k, v in rows:
        print(f"{k}: {v}")
    print(f"\nСохранено: {OUT_FILE}")

if __name__ == "__main__":
    main()
